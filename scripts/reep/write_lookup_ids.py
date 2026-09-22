#!/usr/bin/env python3
"""Write Reep identity columns onto the 'Lookup' sheet of Champions League-201516.xlsx without letting
openpyxl rewrite the workbook (surgical zip/XML edit, the method in reference_surgical_xlsx_cell_edit).

Adds four columns after the sheet's last used column (AE):
    AF  Reep ID        reep_v1_id
    AG  Wikidata QID   wikidata_qid (Reep v0)
    AH  ESPN ID        espn_id
    AI  UEFA ID        uefa_id
Values are written as inline strings (t="inlineStr"), so xl/sharedStrings.xml is never touched. Only
xl/worksheets/sheet15.xml changes; every other zip entry is copied through with its original ZipInfo.

Dry run by default: reports what would change. --write backs the master up beside itself
(.bak-<date>-reepids), writes a .new file, verifies it (testzip, entry list and order, only the Lookup
sheet differs, cells read back), then swaps it in. Refuses if any target cell already holds anything.

    python3 scripts/reep/write_lookup_ids.py scripts/reep/dryrun-2026-09-22-v1/football_team_reep.json
    python3 scripts/reep/write_lookup_ids.py ... --write
"""
import argparse, datetime, json, os, re, shutil, sys, zipfile
from pathlib import Path
from xml.sax.saxutils import escape

MASTER = Path(os.environ.get("CL_WORKBOOK", os.path.expanduser("~/mnt/Excel Files/Champions League-201516.xlsx")))
NEW_COLS = [("AF", "Reep ID", "reep_v1_id"), ("AG", "Wikidata QID", "wikidata_qid"), ("AH", "ESPN ID", "espn_id"), ("AI", "UEFA ID", "uefa_id")]
HEADER_STYLE, CELL_STYLE = "13", "1"

def sheet_path(z):
    wb = z.read("xl/workbook.xml").decode(); rels = z.read("xl/_rels/workbook.xml.rels").decode()
    rid = re.search(r'<sheet [^>]*name="Lookup"[^>]*r:id="(rId\d+)"', wb).group(1)
    tgt = re.search(r'<Relationship [^>]*Id="%s"[^>]*Target="([^"]+)"' % rid, rels) or re.search(r'<Relationship [^>]*Target="([^"]+)"[^>]*Id="%s"' % rid, rels)
    t = tgt.group(1); return t[1:] if t.startswith("/") else "xl/" + t

def cell(ref, val, style):
    return f'<c r="{ref}" s="{style}" t="inlineStr"><is><t>{escape(str(val))}</t></is></c>'

def patch(xml, rows_by_num):
    """rows_by_num: {excel_row: {col_letter: value}}. Appends cells to each <row>, bumps spans and dimension."""
    written = 0; touched = 0
    def fix_row(m):
        nonlocal written, touched
        head, body, tail = m.group(1), m.group(2), m.group(3)
        rnum = int(re.search(r' r="(\d+)"', head).group(1))
        vals = rows_by_num.get(rnum)
        if not vals: return m.group(0)
        for col in vals:
            if re.search(r'<c r="%s%d"' % (col, rnum), body): raise SystemExit(f"cell {col}{rnum} already exists; refusing")
        style = HEADER_STYLE if rnum == 1 else CELL_STYLE
        add = "".join(cell(f"{c}{rnum}", v, style) for c, v in vals.items() if v not in (None, ""))
        if not add: return m.group(0)
        written += sum(1 for v in vals.values() if v not in (None, "")); touched += 1
        head2 = re.sub(r' spans="(\d+):(\d+)"', lambda s: f' spans="{s.group(1)}:{max(int(s.group(2)), 35)}"', head)
        return head2 + body + add + tail
    out = re.sub(r'(<row [^>]*>)(.*?)(</row>)', fix_row, xml, flags=re.S)
    out = re.sub(r'<row ([^>]*)/>', lambda m: m.group(0), out)   # self-closing empty rows carry no cells to extend
    out, n = re.subn(r'<dimension ref="A1:AE(\d+)"/>', r'<dimension ref="A1:AI\1"/>', out)
    if n != 1: raise SystemExit("dimension ref not A1:AE<n>; sheet layout changed, refusing")
    return out, written, touched

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("json_path"); ap.add_argument("--write", action="store_true"); ap.add_argument("--master", default=str(MASTER))
    a = ap.parse_args(); master = Path(a.master)
    rows = json.load(open(a.json_path, encoding="utf-8"))
    plan = {1: {c: h for c, h, _ in NEW_COLS}}
    for r in rows:
        vals = {c: r.get(k) for c, _, k in NEW_COLS}
        if any(vals.values()): plan[int(r["sheet_row"])] = vals
    z = zipfile.ZipFile(master); sp = sheet_path(z)
    xml = z.read(sp).decode("utf-8")
    # sanity: the sheet_row -> team alignment must hold for a sample before anything is written
    import openpyxl
    wb = openpyxl.load_workbook(master, read_only=True, data_only=True); ws = wb["Lookup"]
    sample = [r for r in rows if r["reep_v1_id"]][::700][:12]
    for r in sample:
        v = ws.cell(row=int(r["sheet_row"]), column=1).value
        if v != r["team"]: raise SystemExit(f"row {r['sheet_row']}: workbook says {v!r}, bridge says {r['team']!r}; sheet moved, refusing")
    wb.close()
    new_xml, written, touched = patch(xml, plan)
    print(f"Lookup = {sp}; rows to touch {touched} (incl. header), cells to write {written}; sample alignment OK ({len(sample)} rows)")
    if not a.write: print("dry run; add --write"); return
    stamp = datetime.datetime.now().strftime("%Y%m%d")
    bak = master.with_name(master.name + f".bak-{stamp}-reepids"); new = master.with_name(master.stem + ".reepids-new.xlsx")
    import hashlib
    def h(p):
        m = hashlib.sha256()
        with open(p, "rb") as f:
            for chunk in iter(lambda: f.read(1 << 20), b""): m.update(chunk)
        return m.hexdigest()
    if bak.exists():
        if h(bak) != h(master): raise SystemExit(f"{bak} exists and differs from the master; refusing to overwrite a backup")
        print("backup already present and identical:", bak.name)
    else:
        shutil.copyfile(master, bak); print("backup:", bak.name)
    with zipfile.ZipFile(new, "w") as zo:
        for info in z.infolist():
            data = new_xml.encode("utf-8") if info.filename == sp else z.read(info.filename)
            zo.writestr(info, data, compress_type=info.compress_type)
    z.close()
    # verify
    zn = zipfile.ZipFile(new); zb = zipfile.ZipFile(bak)
    assert zn.testzip() is None, "testzip failed"
    assert zn.namelist() == zb.namelist(), "entry list/order differs"
    diff = [n for n in zb.namelist() if zn.read(n) != zb.read(n)]
    assert diff == [sp], f"unexpected entries differ: {diff}"
    zn.close(); zb.close()
    wb = openpyxl.load_workbook(new, read_only=True, data_only=True); ws = wb["Lookup"]
    assert [ws.cell(row=1, column=32 + i).value for i in range(4)] == [h for _, h, _ in NEW_COLS], "header read-back failed"
    n_ids = 0
    for r in sample:
        assert ws.cell(row=int(r["sheet_row"]), column=32).value == r["reep_v1_id"], f"read-back mismatch at row {r['sheet_row']}"
    for row in ws.iter_rows(min_row=2, min_col=32, max_col=32, values_only=True):
        if row[0]: n_ids += 1
    wb.close()
    expected = sum(1 for r in rows if r["reep_v1_id"])
    assert n_ids == expected, f"Reep ID count {n_ids} != {expected}"
    os.replace(new, master)
    print(f"written: {written} cells on {touched} rows; {n_ids} Reep IDs read back; {master.name} replaced (backup kept)")

if __name__ == "__main__": main()
