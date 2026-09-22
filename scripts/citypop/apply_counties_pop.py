#!/usr/bin/env python3
"""Apply a citypopulation.de population plan to MetroAreas.xlsx 'Counties' by surgical zip/XML edit.

Plan: JSON list of {country, row, name, state, old, new} produced by the diff step (every row already
matched exactly by name and parent; nothing here matches names). For each row it rewrites the Pop cell
(column F) in place, refusing if the cell is a formula or its current value is not `old`, and sets
Census Year (column K) to --census-year, inserting the cell in column order when the row has none.
Only the Counties sheet XML changes; every other zip entry is copied through with its ZipInfo.

Dry run by default. --write backs the master up beside itself (.bak-<date>-<tag>), writes a temp
workbook, verifies (testzip, entry list and order, only the one sheet differs, every planned cell reads
back), then swaps it in.

    python3 scripts/citypop/apply_counties_pop.py <plan.json> --tag thailand-peru --census-year 2025
    python3 scripts/citypop/apply_counties_pop.py <plan.json> --tag thailand-peru --census-year 2025 --write
"""
import argparse, datetime, hashlib, json, os, re, shutil, zipfile
from pathlib import Path

MASTER = Path(os.environ.get("METRO_WORKBOOK", os.path.expanduser("~/mnt/Excel Files/MetroAreas.xlsx")))
POP_COL, CENSUS_COL = "F", "K"

def col_index(letters):
    n = 0
    for ch in letters: n = n * 26 + (ord(ch) - 64)
    return n

def sheet_path(z, name):
    wb = z.read("xl/workbook.xml").decode(); rels = z.read("xl/_rels/workbook.xml.rels").decode()
    rid = re.search(r'<sheet [^>]*name="%s"[^>]*r:id="(rId\d+)"' % re.escape(name), wb).group(1)
    tgt = re.search(r'<Relationship [^>]*Id="%s"[^>]*Target="([^"]+)"' % rid, rels) or re.search(r'<Relationship [^>]*Target="([^"]+)"[^>]*Id="%s"' % rid, rels)
    t = tgt.group(1); return t[1:] if t.startswith("/") else "xl/" + t

def patch(xml, plan, census_year):
    by_row = {int(p["row"]): p for p in plan}
    done = {"pop": 0, "census_set": 0, "census_added": 0}
    def fix_row(m):
        head, body, tail = m.group(1), m.group(2), m.group(3)
        rnum = int(re.search(r' r="(\d+)"', head).group(1))
        p = by_row.get(rnum)
        if not p: return m.group(0)
        # Pop cell: must exist, be numeric, hold the expected old value
        cm = re.search(r'<c r="%s%d"([^>]*)>(.*?)</c>' % (POP_COL, rnum), body, re.S)
        if not cm: raise SystemExit(f"row {rnum} {p['name']}: no Pop cell")
        attrs, inner = cm.group(1), cm.group(2)
        if "<f" in inner or 't="s"' in attrs or 't="str"' in attrs: raise SystemExit(f"row {rnum} {p['name']}: Pop cell is a formula or text; refusing")
        vm = re.search(r"<v>([^<]*)</v>", inner)
        cur = float(vm.group(1)) if vm else None
        if cur is None or abs(cur - float(p["old"])) > 0.5: raise SystemExit(f"row {rnum} {p['name']}: Pop is {cur}, plan expected {p['old']}; refusing")
        new_cell = f'<c r="{POP_COL}{rnum}"{attrs}><v>{int(p["new"])}</v></c>'
        body = body[:cm.start()] + new_cell + body[cm.end():]; done["pop"] += 1
        # Census Year cell: rewrite or insert in column order
        km = re.search(r'<c r="%s%d"([^>]*?)(/>|>(.*?)</c>)' % (CENSUS_COL, rnum), body, re.S)
        if km:
            kattrs = km.group(1)
            if km.group(3) and "<f" in km.group(3): raise SystemExit(f"row {rnum}: Census Year is a formula; refusing")
            kattrs = re.sub(r' t="[^"]*"', "", kattrs)
            body = body[:km.start()] + f'<c r="{CENSUS_COL}{rnum}"{kattrs}><v>{census_year}</v></c>' + body[km.end():]; done["census_set"] += 1
        else:
            # style: reuse the Pop cell's style attribute if it has one
            sm = re.search(r' s="\d+"', attrs); style = sm.group(0) if sm else ""
            cell = f'<c r="{CENSUS_COL}{rnum}"{style}><v>{census_year}</v></c>'
            pos = len(body)
            for c in re.finditer(r'<c r="([A-Z]+)%d"' % rnum, body):
                if col_index(c.group(1)) > col_index(CENSUS_COL): pos = c.start(); break
            body = body[:pos] + cell + body[pos:]; done["census_added"] += 1
        return head + body + tail
    out = re.sub(r'(<row [^>]*>)(.*?)(</row>)', fix_row, xml, flags=re.S)
    return out, done

def sha(p):
    m = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""): m.update(chunk)
    return m.hexdigest()

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("plan"); ap.add_argument("--tag", required=True); ap.add_argument("--census-year", type=int, required=True)
    ap.add_argument("--write", action="store_true"); ap.add_argument("--master", default=str(MASTER)); ap.add_argument("--sheet", default="Counties")
    a = ap.parse_args(); master = Path(a.master); plan = json.load(open(a.plan, encoding="utf-8"))
    z = zipfile.ZipFile(master); sp = sheet_path(z, a.sheet); xml = z.read(sp).decode("utf-8")
    new_xml, done = patch(xml, plan, a.census_year)
    if done["pop"] != len(plan): raise SystemExit(f"planned {len(plan)} rows, patched {done['pop']}; refusing")
    print(f"{a.sheet} = {sp}; plan rows {len(plan)}; Pop cells to rewrite {done['pop']}; Census Year set {done['census_set']} / added {done['census_added']}")
    if not a.write: print("dry run; add --write"); return
    stamp = datetime.datetime.now().strftime("%Y%m%d")
    bak = master.with_name(master.name + f".bak-{stamp}-{a.tag}"); new = master.with_name(master.stem + f".{a.tag}-new.xlsx")
    if bak.exists():
        if sha(bak) != sha(master): raise SystemExit(f"{bak.name} exists and differs from the master; refusing")
        print("backup already present and identical:", bak.name)
    else:
        shutil.copyfile(master, bak); print("backup:", bak.name)
    with zipfile.ZipFile(new, "w") as zo:
        for info in z.infolist():
            zo.writestr(info, new_xml.encode("utf-8") if info.filename == sp else z.read(info.filename), compress_type=info.compress_type)
    z.close()
    zn = zipfile.ZipFile(new); zb = zipfile.ZipFile(bak)
    assert zn.testzip() is None, "testzip failed"
    assert zn.namelist() == zb.namelist(), "entry list/order differs"
    diff = [n for n in zb.namelist() if zn.read(n) != zb.read(n)]
    assert diff == [sp], f"unexpected entries differ: {diff}"
    zn.close(); zb.close()
    import openpyxl
    wb = openpyxl.load_workbook(new, read_only=True, data_only=True); ws = wb[a.sheet]
    bad = 0
    for p in plan:
        r = int(p["row"])
        if ws.cell(row=r, column=col_index(POP_COL)).value != int(p["new"]) or ws.cell(row=r, column=col_index(CENSUS_COL)).value != a.census_year or ws.cell(row=r, column=2).value != p["name"]: bad += 1
    wb.close()
    assert bad == 0, f"{bad} rows failed read-back"
    os.replace(new, master)
    print(f"written: {len(plan)} rows; all read back; {master.name} replaced (backup kept)")

if __name__ == "__main__": main()
